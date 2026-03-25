#!/usr/bin/env python3
"""NemoClaw/OpenClaw 構成図の PowerPoint を生成する"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# ── 色定義 ──
BG_COLOR = RGBColor(0x1A, 0x1A, 0x2E)
NVIDIA_GREEN = RGBColor(0x76, 0xB9, 0x00)
DISCORD_PURPLE = RGBColor(0x5B, 0x65, 0xEA)
DOCKER_BLUE = RGBColor(0x00, 0x97, 0xE6)
SAND_ORANGE = RGBColor(0xFF, 0x8C, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY = RGBColor(0xBB, 0xBB, 0xBB)
DARK_GRAY = RGBColor(0x2A, 0x2A, 0x40)
ARROW_CYAN = RGBColor(0x00, 0xD4, 0xFF)
CARD_BG = RGBColor(0x25, 0x25, 0x3D)


def set_slide_bg(slide, color):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_box(slide, left, top, width, height, fill_color, border_color=None,
            text="", font_size=14, font_color=WHITE, bold=False, align=PP_ALIGN.CENTER):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(2)
    else:
        shape.line.fill.background()
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = font_color
    p.font.bold = bold
    p.alignment = align
    return shape


def add_text(slide, left, top, width, height, text, font_size=12,
             font_color=WHITE, bold=False, align=PP_ALIGN.LEFT):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = font_color
    p.font.bold = bold
    p.alignment = align
    return txBox


def add_arrow_line(slide, x1, y1, x2, y2, color=ARROW_CYAN, width=Pt(3)):
    connector = slide.shapes.add_connector(1, x1, y1, x2, y2)  # straight
    connector.line.color.rgb = color
    connector.line.width = width


# ═══════════════════════════════════════════════════
# Slide 1: タイトル
# ═══════════════════════════════════════════════════
slide1 = prs.slides.add_slide(prs.slide_layouts[6])  # blank
set_slide_bg(slide1, BG_COLOR)

add_text(slide1, Inches(1), Inches(1.5), Inches(11), Inches(1.2),
         "NemoClaw / OpenClaw", font_size=44, font_color=NVIDIA_GREEN, bold=True,
         align=PP_ALIGN.CENTER)
add_text(slide1, Inches(1), Inches(2.8), Inches(11), Inches(0.8),
         "Discord Bridge 構成図", font_size=28, font_color=WHITE, bold=False,
         align=PP_ALIGN.CENTER)
add_text(slide1, Inches(1), Inches(4.0), Inches(11), Inches(0.6),
         "Surface Laptop 3  |  Windows + WSL2 (Ubuntu)  |  2026-03-22",
         font_size=16, font_color=LIGHT_GRAY, align=PP_ALIGN.CENTER)

# ═══════════════════════════════════════════════════
# Slide 2: 全体アーキテクチャ図
# ═══════════════════════════════════════════════════
slide2 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide2, BG_COLOR)

add_text(slide2, Inches(0.5), Inches(0.2), Inches(12), Inches(0.6),
         "全体アーキテクチャ", font_size=28, font_color=NVIDIA_GREEN, bold=True)

# -- Discord Cloud --
add_box(slide2, Inches(0.5), Inches(1.5), Inches(2.2), Inches(1.2),
        DISCORD_PURPLE, border_color=DISCORD_PURPLE,
        text="Discord\n(Cloud)", font_size=16, bold=True)

# -- Arrow: Discord → Bot --
add_arrow_line(slide2, Inches(2.7), Inches(2.1), Inches(3.5), Inches(2.1))

# -- Discord Bridge Bot (WSL) --
add_box(slide2, Inches(3.5), Inches(1.2), Inches(2.5), Inches(1.8),
        CARD_BG, border_color=DOCKER_BLUE,
        text="Discord Bridge Bot\n(Python / WSL2)\ndiscord.py + asyncio",
        font_size=13, bold=True)

# -- Arrow: Bot → SSH --
add_arrow_line(slide2, Inches(6.0), Inches(2.1), Inches(6.8), Inches(2.1))

# -- SSH proxy label --
add_text(slide2, Inches(6.1), Inches(1.4), Inches(1.5), Inches(0.5),
         "SSH", font_size=11, font_color=ARROW_CYAN, align=PP_ALIGN.CENTER)
add_text(slide2, Inches(6.1), Inches(1.7), Inches(1.5), Inches(0.5),
         "(openshell\nssh-proxy)", font_size=9, font_color=LIGHT_GRAY,
         align=PP_ALIGN.CENTER)

# -- OpenShell Gateway (k3s) --
add_box(slide2, Inches(6.8), Inches(0.9), Inches(5.5), Inches(3.0),
        DARK_GRAY, border_color=NVIDIA_GREEN,
        text="", font_size=1)

add_text(slide2, Inches(7.0), Inches(0.95), Inches(5.0), Inches(0.4),
         "OpenShell Gateway (k3s Docker)", font_size=14,
         font_color=NVIDIA_GREEN, bold=True)

# -- Sandbox --
add_box(slide2, Inches(7.2), Inches(1.5), Inches(4.8), Inches(2.1),
        RGBColor(0x1E, 0x1E, 0x35), border_color=SAND_ORANGE,
        text="", font_size=1)

add_text(slide2, Inches(7.3), Inches(1.55), Inches(4.5), Inches(0.3),
         "Sandbox: my-assistant  (Landlock + seccomp + netns)",
         font_size=11, font_color=SAND_ORANGE, bold=True)

# -- OpenClaw Agent --
add_box(slide2, Inches(7.5), Inches(2.0), Inches(2.0), Inches(1.3),
        CARD_BG, border_color=NVIDIA_GREEN,
        text="OpenClaw Agent\n(main)\nv2026.3.11", font_size=12, bold=True)

# -- /sandbox workspace --
add_box(slide2, Inches(9.8), Inches(2.0), Inches(1.9), Inches(1.3),
        CARD_BG, border_color=LIGHT_GRAY,
        text="/sandbox/\nworkspace\nUSER.md", font_size=11)

# -- Arrow: Agent → NVIDIA Cloud --
add_arrow_line(slide2, Inches(8.5), Inches(3.5), Inches(8.5), Inches(4.5))

# -- NVIDIA Cloud API --
add_box(slide2, Inches(7.0), Inches(4.5), Inches(3.2), Inches(1.2),
        RGBColor(0x1A, 0x40, 0x1A), border_color=NVIDIA_GREEN,
        text="NVIDIA Cloud API\nnemotron-3-super-120b-a12b\n(inference.local)",
        font_size=12, bold=True)

# -- Security labels --
add_text(slide2, Inches(0.5), Inches(4.8), Inches(5.5), Inches(1.5),
         "セキュリティ境界:\n"
         "• Landlock: ファイルシステム隔離 (/sandbox, /tmp のみ RW)\n"
         "• seccomp: システムコール制限\n"
         "• netns: ネットワーク名前空間隔離\n"
         "• ポリシー: pypi, npm, discord のみ許可",
         font_size=11, font_color=LIGHT_GRAY)

# ═══════════════════════════════════════════════════
# Slide 3: データフロー
# ═══════════════════════════════════════════════════
slide3 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide3, BG_COLOR)

add_text(slide3, Inches(0.5), Inches(0.2), Inches(12), Inches(0.6),
         "データフロー (メッセージ処理)", font_size=28,
         font_color=NVIDIA_GREEN, bold=True)

steps = [
    ("1", "ユーザーが\nDiscord に投稿", DISCORD_PURPLE),
    ("2", "Bot がメッセージ\nを受信 (discord.py)", DOCKER_BLUE),
    ("3", "SSH 経由で\nサンドボックスに転送", ARROW_CYAN),
    ("4", "openclaw agent\n-m でメッセージ送信", NVIDIA_GREEN),
    ("5", "NVIDIA Cloud API\nで推論実行", RGBColor(0x1A, 0x80, 0x1A)),
    ("6", "応答をクリーニング\nして Discord 返信", DISCORD_PURPLE),
]

for i, (num, label, color) in enumerate(steps):
    x = Inches(0.4 + i * 2.1)
    y = Inches(1.5)
    add_box(slide3, x, y, Inches(1.9), Inches(1.6), CARD_BG,
            border_color=color, text=f"{num}\n{label}", font_size=12, bold=True)
    if i < len(steps) - 1:
        add_arrow_line(slide3, x + Inches(1.9), y + Inches(0.8),
                       x + Inches(2.1), y + Inches(0.8), color=ARROW_CYAN)

# key details
details_text = (
    "ポイント:\n"
    "• SSH 引数は shlex.quote() でエスケープ（スペース / 日本語対応）\n"
    "• clean_response() で [plugins] / ANSI / 連続改行を除去\n"
    "• タイムアウト: 180秒  |  最大応答: 1200文字  |  Discord制限: 2000文字\n"
    "• セッションID: discord-{channel_id} でチャンネルごとに会話維持"
)
add_text(slide3, Inches(0.5), Inches(3.8), Inches(12), Inches(2.5),
         details_text, font_size=13, font_color=LIGHT_GRAY)

# ═══════════════════════════════════════════════════
# Slide 4: コンポーネント詳細
# ═══════════════════════════════════════════════════
slide4 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide4, BG_COLOR)

add_text(slide4, Inches(0.5), Inches(0.2), Inches(12), Inches(0.6),
         "コンポーネント詳細", font_size=28, font_color=NVIDIA_GREEN, bold=True)

components = [
    ("Discord Bridge Bot", DISCORD_PURPLE, [
        "言語: Python 3.12 + discord.py",
        "場所: /tmp/discord-bot/bot.py (WSL)",
        "起動: start-discord-bridge.sh",
        "全メッセージ応答モード (RESPOND_TO_ALL=true)",
        "日本語 + 簡潔回答プロンプト自動付与",
    ]),
    ("OpenShell Gateway", DOCKER_BLUE, [
        "コンテナ: openshell-cluster-nemoclaw",
        "Docker CE 29.3.0 (k3s 内部)",
        "SSH ProxyCommand: openshell ssh-proxy",
        "ポート: 8080 (Gateway)",
        "MTU: 1350 (WSL2 TLS対策)",
    ]),
    ("OpenClaw Sandbox", SAND_ORANGE, [
        "名前: my-assistant",
        "イメージ: openclaw:latest",
        "ポリシー: Landlock + seccomp + netns",
        "RW: /sandbox, /tmp",
        "RO: /usr, /lib, /app, /etc",
    ]),
    ("NVIDIA Inference", NVIDIA_GREEN, [
        "Model: nemotron-3-super-120b-a12b",
        "Provider: NVIDIA Cloud API",
        "Endpoint: inference.local (managed)",
        "Plugin: NemoClaw",
        "Profile: default (NCP対応)",
    ]),
]

for i, (title, color, items) in enumerate(components):
    x = Inches(0.3 + i * 3.2)
    add_box(slide4, x, Inches(1.0), Inches(3.0), Inches(0.6), color,
            text=title, font_size=14, bold=True)
    body = "\n".join(f"• {item}" for item in items)
    add_box(slide4, x, Inches(1.6), Inches(3.0), Inches(3.5), CARD_BG,
            border_color=color, text=body, font_size=11,
            align=PP_ALIGN.LEFT)

# ═══════════════════════════════════════════════════
# Slide 5: 運用コマンド
# ═══════════════════════════════════════════════════
slide5 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide5, BG_COLOR)

add_text(slide5, Inches(0.5), Inches(0.2), Inches(12), Inches(0.6),
         "運用コマンド", font_size=28, font_color=NVIDIA_GREEN, bold=True)

commands = [
    ("WSL 再起動後の復旧", [
        "sudo service docker start",
        "sudo ip link set dev eth0 mtu 1350",
        "sudo sysctl -w net.ipv4.tcp_mtu_probing=1",
    ]),
    ("OpenClaw 起動", [
        "source ~/.nvm/nvm.sh",
        "nemoclaw my-assistant status",
        "nemoclaw my-assistant connect",
        "openclaw tui",
    ]),
    ("Discord Bot 管理", [
        "bash start-discord-bridge.sh",
        "bash stop-discord-bridge.sh",
        "tail -f /tmp/discord-bot/bot.log",
    ]),
    ("トラブルシューティング", [
        "docker rm -f openshell-cluster-nemoclaw",
        "openshell gateway destroy --name nemoclaw",
        "nemoclaw onboard  # 完全リセット",
    ]),
]

for i, (title, cmds) in enumerate(commands):
    y = Inches(1.0 + i * 1.5)
    add_text(slide5, Inches(0.5), y, Inches(3.5), Inches(0.4),
             title, font_size=15, font_color=NVIDIA_GREEN, bold=True)
    cmd_text = "\n".join(f"$ {c}" for c in cmds)
    add_box(slide5, Inches(4.2), y, Inches(8.5), Inches(1.2),
            DARK_GRAY, border_color=LIGHT_GRAY,
            text=cmd_text, font_size=11, font_color=RGBColor(0x80, 0xFF, 0x80),
            align=PP_ALIGN.LEFT)

# ── 保存 ──
out_path = "/mnt/c/Users/kazuk/NemoClaw/NemoClaw-Architecture.pptx"
prs.save(out_path)
print(f"Saved: {out_path}")
